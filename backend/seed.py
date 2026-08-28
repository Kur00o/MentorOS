"""
Seed the local dev database from the real MITWPU CSE mentor-allocation
spreadsheet ("Mentor Data.xlsx").

Mode B (default): REAL mentors, ANONYMIZED students — mentor names/emails are
kept (needed for login + rosters), but every student's name/email/phone is
replaced with a synthetic value so no student PII lands in the dev DB.

The spreadsheet is a block-structured allocation sheet: a mentor is named on the
first student of their group, and the following blank rows belong to that same
mentor until the next mentor appears (forward-fill).

Usage (run from the project root):
    backend/.venv/Scripts/python.exe -m backend.seed              # anonymized students (mode B)
    backend/.venv/Scripts/python.exe -m backend.seed --real-students   # keep real PII (mode A)
    backend/.venv/Scripts/python.exe -m backend.seed --reset      # wipe seeded rows, then reseed
    backend/.venv/Scripts/python.exe -m backend.seed --tab 23-24  # use a different year tab

Idempotent: re-running upserts by email. --reset clears Student/Mentor/Meeting
rows and their Users (HOD/Admin users are left untouched).
"""
from __future__ import annotations

import argparse
import os
import random
import re
import sys
from datetime import datetime, timedelta, timezone

# Make `backend` importable when run as a script or module.
PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if PROJECT_ROOT not in sys.path:
    sys.path.insert(0, PROJECT_ROOT)

import openpyxl  # noqa: E402

from backend.app.core.database import Base, SessionLocal, engine  # noqa: E402
import backend.app.models  # noqa: F401,E402  (registers all tables on Base)
from backend.app.models.user import User, UserRole  # noqa: E402
from backend.app.models.mentor import Mentor  # noqa: E402
from backend.app.models.student import Student  # noqa: E402
from backend.app.models.meeting import Meeting, MeetingLog  # noqa: E402
from backend.app.core.security import get_password_hash  # noqa: E402

XLSX_PATH = r"D:\=MENTOROS\Mentor Data.xlsx"
DEPARTMENT = "CSE"
SEMESTER = 3  # SY B.Tech = 2nd year, odd semester
EMAIL_RE = re.compile(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}")
PHONE_RE = re.compile(r"(?<!\d)(\d{10,12})(?!\d)")

FIRST_NAMES = [
    "Aarav", "Vivaan", "Aditya", "Vihaan", "Arjun", "Sai", "Reyansh", "Krishna",
    "Ishaan", "Rohan", "Ananya", "Diya", "Aadhya", "Saanvi", "Ira", "Myra",
    "Aarohi", "Anika", "Navya", "Kiara", "Riya", "Neha", "Priya", "Sneha",
    "Kabir", "Dhruv", "Aryan", "Kian", "Yash", "Om", "Atharv", "Shaurya",
]
LAST_NAMES = [
    "Sharma", "Verma", "Patil", "Deshmukh", "Kulkarni", "Joshi", "Gupta",
    "Iyer", "Nair", "Reddy", "Mehta", "Shah", "Rao", "Naik", "Kale", "Jadhav",
    "Chavan", "Pawar", "Bhosale", "Kadam", "More", "Shinde", "Gaikwad", "Wagh",
]


def clean_num(value) -> str:
    """Excel stores numbers as floats (e.g. 9404612559.0); render as a clean int string."""
    if value is None:
        return ""
    if isinstance(value, float):
        return str(int(value))
    return str(value).strip()


def title_from_email(email: str) -> str:
    """devendra.joshi@mitwpu.edu.in -> 'Prof. Devendra Joshi'"""
    local = email.split("@", 1)[0]
    parts = [p for p in re.split(r"[._\-]+", local) if p]
    name = " ".join(p.capitalize() for p in parts)
    return f"Prof. {name}" if name else email


def parse_mentor_cell(cell: str):
    """Extract (email, mobile) of the active mentor from a messy free-text cell."""
    text = str(cell)
    email = None
    for m in EMAIL_RE.findall(text):
        low = m.lower()
        local = low.split("@", 1)[0]
        # Skip numeric student-code emails; require an alphabetic mentor localpart.
        if "mitwpu" in low and re.search(r"[a-z]", local):
            email = low
            break
    mobile_match = PHONE_RE.search(text.replace(" ", ""))
    mobile = mobile_match.group(1) if mobile_match else None
    return email, mobile


def load_allocations(tab: str):
    """Return (mentors_by_email, students) parsed from the given tab."""
    wb = openpyxl.load_workbook(XLSX_PATH, read_only=True, data_only=True)
    if tab not in wb.sheetnames:
        raise SystemExit(f"Tab '{tab}' not found. Available: {wb.sheetnames}")
    ws = wb[tab]

    def g(row, i):
        return row[i] if i < len(row) else None

    rows = list(ws.iter_rows(values_only=True))
    data = rows[2:]  # skip the title + header rows

    mentors: dict[str, dict] = {}
    students: list[dict] = []
    current_email = None

    for r in data:
        name = g(r, 3)
        if not name or not isinstance(name, str) or len(name.strip()) < 2:
            continue
        mentor_cell = g(r, 6)
        if mentor_cell:
            email, mobile = parse_mentor_cell(mentor_cell)
            if email:
                current_email = email
                if email not in mentors:
                    mentors[email] = {
                        "email": email,
                        "name": title_from_email(email),
                        "mobile": mobile,
                    }
                elif mobile and not mentors[email]["mobile"]:
                    mentors[email]["mobile"] = mobile
        if current_email is None:
            continue  # no mentor assigned yet — skip orphan students
        students.append(
            {
                "mentor_email": current_email,
                "real_name": name.strip(),
                "real_email": (clean_num(g(r, 4)) or "").strip().lower() or None,
                "real_mobile": clean_num(g(r, 5)) or None,
                "prn": clean_num(g(r, 1)) or None,
            }
        )
    return mentors, students


def reset_seed(db):
    """Delete seeded data. Leaves HOD/Admin users intact."""
    db.query(MeetingLog).delete()
    db.query(Meeting).delete()
    db.query(Student).delete()
    db.query(Mentor).delete()
    db.query(User).filter(User.role.in_([UserRole.STUDENT, UserRole.MENTOR])).delete(
        synchronize_session=False
    )
    db.commit()


def get_or_create_user(db, email, full_name, role, password):
    user = db.query(User).filter(User.email == email).first()
    if user:
        user.full_name = full_name
        user.role = role
        return user
    user = User(
        email=email,
        full_name=full_name,
        role=role,
        hashed_password=get_password_hash(password),
        is_active=True,
    )
    db.add(user)
    db.flush()
    return user


def seed(tab: str, anonymize: bool, do_reset: bool):
    Base.metadata.create_all(bind=engine)
    rng = random.Random(42)
    now = datetime.now(timezone.utc)

    mentors_by_email, students = load_allocations(tab)
    print(f"Parsed tab '{tab}': {len(mentors_by_email)} mentors, {len(students)} students")

    db = SessionLocal()
    try:
        if do_reset:
            reset_seed(db)
            print("Reset: cleared previously seeded Student/Mentor/Meeting rows.")

        # --- Mentors ---
        mentor_id_by_email: dict[str, int] = {}
        roster_size = {e: sum(1 for s in students if s["mentor_email"] == e) for e in mentors_by_email}
        for email, info in mentors_by_email.items():
            user = get_or_create_user(db, email, info["name"], UserRole.MENTOR, "Mentor@123")
            mentor = db.query(Mentor).filter(Mentor.user_id == user.id).first()
            if not mentor:
                mentor = Mentor(user_id=user.id)
                db.add(mentor)
            mentor.department = DEPARTMENT
            mentor.max_mentees = max(20, roster_size.get(email, 0) + 5)
            mentor.mobile_no = info["mobile"]
            db.flush()
            mentor_id_by_email[email] = mentor.id
        db.commit()

        # --- Students ---
        created = 0
        for idx, s in enumerate(students, start=1):
            if anonymize:
                fn = FIRST_NAMES[idx % len(FIRST_NAMES)]
                ln = LAST_NAMES[(idx // len(FIRST_NAMES)) % len(LAST_NAMES)]
                full_name = f"{fn} {ln}"
                email = f"student{idx:04d}@demo.mentoros.local"
                mobile = f"90{idx:08d}"
                usn = f"1MW24CS{idx:04d}"
            else:
                full_name = s["real_name"]
                email = s["real_email"] or f"student{idx:04d}@demo.mentoros.local"
                mobile = s["real_mobile"]
                usn = s["prn"] or f"1MW24CS{idx:04d}"

            user = get_or_create_user(db, email, full_name, UserRole.STUDENT, "Student@123")
            student = db.query(Student).filter(Student.user_id == user.id).first()
            if not student:
                student = Student(user_id=user.id, usn=usn)
                db.add(student)

            score = round(rng.uniform(35, 95), 1)
            risk = "Green" if score >= 70 else ("Amber" if score >= 50 else "Coral")
            student.usn = usn
            student.department = DEPARTMENT
            student.semester = SEMESTER
            student.student_mobile = mobile
            student.attendance_rate = round(rng.uniform(55, 98), 1)
            student.sgpa = round(rng.uniform(4.5, 9.5), 2)
            student.cgpa = round(min(10.0, max(0.0, student.sgpa + rng.uniform(-0.6, 0.6))), 2)
            student.success_score = score
            student.risk_status = risk
            student.consent_given = True
            student.is_under_18 = False
            student.mentor_id = mentor_id_by_email[s["mentor_email"]]
            db.flush()
            created += 1
        db.commit()

        # --- Meetings (only if none seeded yet, keeps it idempotent) ---
        meetings_made = 0
        if db.query(Meeting).count() == 0:
            all_students = db.query(Student).all()
            for st in all_students:
                roll = rng.random()
                if roll < 0.45:  # a past, completed meeting
                    m = Meeting(
                        mentor_id=st.mentor_id,
                        student_id=st.id,
                        title="Mid-semester check-in",
                        date=now - timedelta(days=rng.randint(5, 40)),
                        mode=rng.choice(["in-person", "video"]),
                        status="Completed",
                    )
                    db.add(m)
                    db.flush()
                    if rng.random() < 0.5:
                        db.add(MeetingLog(
                            meeting_id=m.id,
                            topics_discussed=["Academic progress", "Attendance"],
                            action_items=["Submit pending assignment"],
                            observations="Discussed improvement plan.",
                            logged_by=st.mentor_id and db.query(Mentor).get(st.mentor_id).user_id,
                        ))
                    meetings_made += 1
                elif roll < 0.75:  # an upcoming scheduled meeting
                    db.add(Meeting(
                        mentor_id=st.mentor_id,
                        student_id=st.id,
                        title="Progress review",
                        date=now + timedelta(days=rng.randint(2, 21)),
                        mode=rng.choice(["in-person", "video"]),
                        status="Scheduled",
                    ))
                    meetings_made += 1
            db.commit()

        # --- Summary ---
        print("\n=== Seed complete ===")
        print(f"Mode:            {'B (anonymized students)' if anonymize else 'A (real student PII)'}")
        print(f"Mentors:         {db.query(Mentor).count()}")
        print(f"Students:        {db.query(Student).count()}")
        print(f"Meetings:        {db.query(Meeting).count()}")
        print(f"Meeting logs:    {db.query(MeetingLog).count()}")
        sample = db.query(Mentor).first()
        if sample:
            n = db.query(Student).filter(Student.mentor_id == sample.id).count()
            print(f"Sample mentor:   {sample.user.email}  ({n} mentees)  login pw: Mentor@123")
    finally:
        db.close()


def main():
    ap = argparse.ArgumentParser(description="Seed the MentorOS dev DB from the allocation spreadsheet.")
    ap.add_argument("--tab", default="24-25", help="Worksheet/year tab to load (default: 24-25)")
    ap.add_argument("--real-students", action="store_true", help="Keep real student PII (mode A)")
    ap.add_argument("--reset", action="store_true", help="Clear seeded rows before seeding")
    args = ap.parse_args()
    seed(tab=args.tab, anonymize=not args.real_students, do_reset=args.reset)


if __name__ == "__main__":
    main()
